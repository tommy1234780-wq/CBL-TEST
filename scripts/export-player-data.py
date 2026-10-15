#!/usr/bin/env python3
"""Export CBL's official player workbook into website-safe JSON.

The website searches by player name, but joins season and career records with the
permanent player ID. Only approved public fields are exported.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import defaultdict
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


DEFAULT_SOURCE = Path(
    r"D:\定銨\CBL\Mick\第三屆\裁判組紀錄組賽程組\CBL歷屆球員成績資料庫_正式版.xlsx"
)
DEFAULT_HITTER = DEFAULT_SOURCE.with_name("CBL第三屆打者個人成績登錄系統_正式版.xlsx")
DEFAULT_PITCHER = DEFAULT_SOURCE.with_name("CBL第三屆投手個人成績登錄系統_正式版.xlsx")

HITTER_HEADERS = [
    "屆數", "隊伍", "姓名", "組別", "球員ID", "出賽紀錄", "打席", "打數", "打擊率",
    "一壘打", "二壘打", "三壘打", "全壘打", "安打", "三振", "保送", "失誤上壘",
    "觸身球", "內滾出", "飛球出", "野選", "妨礙守備出局", "其他打數", "犧牲觸擊",
    "犧牲飛球", "妨礙上壘", "其他打席", "打點", "雙殺打", "盜壘", "壘打數",
    "守備失誤", "捕手阻殺", "捕逸", "被盜壘", "上壘率", "長打率", "被三振率", "OPS",
]
PITCHER_HEADERS = [
    "屆數", "隊伍", "姓名", "組別", "球員ID", "出賽紀錄", "投球數", "先發", "中繼成功",
    "救援成功", "勝投", "敗投", "總出局數", "投球局數", "總失分", "自責分", "非自責分",
    "三振", "保送", "觸身", "被安打", "被全壘打", "失誤", "暴投", "投手犯規", "防禦率",
    "K9", "WHIP", "BB9", "H9",
]
HITTER_DERIVED = {"打擊率", "上壘率", "長打率", "被三振率", "OPS"}
PITCHER_DERIVED = {"投球局數", "非自責分", "防禦率", "K9", "WHIP", "BB9", "H9"}
HITTER_ALIASES = {"妨礙上壘": ["妨礙上壘", "妨礙打擊上壘"]}
PUBLIC_INTERNAL_KEYS = {"姓名別名"}
TEAM_NAME_ALIASES = {
    "好丘": "好CHILL",
    "好丘棒球隊": "好CHILL",
    "好chill": "好CHILL",
    "好chill棒球隊": "好CHILL",
}


def parse_args() -> argparse.Namespace:
    script_root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(description="Export official CBL player data for the website")
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--hitter-workbook", type=Path, default=DEFAULT_HITTER)
    parser.add_argument("--pitcher-workbook", type=Path, default=DEFAULT_PITCHER)
    parser.add_argument("--output", type=Path, default=script_root / "assets" / "data" / "player-stats.json")
    return parser.parse_args()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def clean(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return "" if value is None else value


def table_rows(sheet: Any, header_row: int = 4) -> list[dict[str, Any]]:
    headers = [clean(cell.value) for cell in sheet[header_row]]
    rows: list[dict[str, Any]] = []
    for values in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        row = {str(header): clean(values[index] if index < len(values) else "") for index, header in enumerate(headers) if header}
        if any(value not in (None, "") for value in row.values()):
            rows.append(row)
    return rows


def split_aliases(value: Any) -> list[str]:
    return [part.strip() for part in re.split(r"[、,，;；/|\n]+", str(value or "")) if part.strip()]


def canonical_team_name(value: Any) -> str:
    original = str(value or "").strip()
    key = re.sub(r"[\s.．・·_-]+", "", original).casefold()
    return TEAM_NAME_ALIASES.get(key, original)


def number(value: Any) -> float:
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def compact_number(value: float) -> int | float:
    rounded = round(value, 10)
    return int(rounded) if float(rounded).is_integer() else rounded


def rate(numerator: float, denominator: float, digits: int) -> float:
    return round(numerator / denominator, digits) if denominator else 0.0


def innings_from_outs(outs: float) -> str:
    whole = int(outs // 3)
    remainder = int(round(outs - whole * 3))
    return f"{whole}.{remainder}"


def approved_historical_row(row: dict[str, Any]) -> bool:
    return (
        row.get("球員ID")
        and row.get("資料檢核") == "OK"
        and row.get("納入狀態") == "可納入"
        and row.get("身分確認") == "已確認"
        and row.get("重複檢核") == "OK"
    )


def public_row(
    row: dict[str, Any],
    headers: list[str],
    player: dict[str, Any],
    aliases: dict[str, list[str]] | None = None,
) -> dict[str, Any]:
    aliases = aliases or {}
    output: dict[str, Any] = {}
    for header in headers:
        candidates = aliases.get(header, [header])
        output[header] = next((clean(row.get(candidate)) for candidate in candidates if row.get(candidate) not in (None, "")), "")
        if header == "隊伍":
            output[header] = canonical_team_name(output[header])
        if output[header] != "":
            if header in HITTER_DERIVED:
                output[header] = round(number(output[header]), 3)
            elif header in PITCHER_DERIVED - {"投球局數", "非自責分"}:
                output[header] = round(number(output[header]), 2)
    output["球員ID"] = player["id"]
    output["姓名"] = player["name"]
    output["姓名別名"] = "、".join(player["aliases"])
    return output


def meaningful_settled_row(row: dict[str, Any], stat_headers: Iterable[str]) -> bool:
    if not row.get("球員ID") or row.get("資料檢核") != "OK":
        return False
    if row.get("最後結算批次") or row.get("最後結算日期"):
        return True
    return any(number(row.get(header)) != 0 for header in stat_headers)


def add_third_season_rows(
    workbook_path: Path,
    headers: list[str],
    players: dict[str, dict[str, Any]],
    rows: list[dict[str, Any]],
    aliases: dict[str, list[str]] | None = None,
) -> None:
    if not workbook_path.exists():
        return
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    settled_rows = table_rows(workbook["已結算累計（系統）"])
    stat_headers = [header for header in headers if header not in {"屆數", "隊伍", "姓名", "組別", "球員ID"}]
    existing = {(row["屆數"], row["隊伍"], row["球員ID"]) for row in rows}
    for row in settled_rows:
        player_id = str(row.get("球員ID") or "").strip()
        if player_id not in players or not meaningful_settled_row(row, stat_headers):
            continue
        exported = public_row(row, headers, players[player_id], aliases)
        key = (str(exported.get("屆數") or "第三屆"), str(exported.get("隊伍") or ""), player_id)
        if key in existing:
            continue
        rows.append(exported)
        existing.add(key)


def career_rows(
    rows: list[dict[str, Any]],
    headers: list[str],
    players: dict[str, dict[str, Any]],
    derived_headers: set[str],
    kind: str,
) -> list[dict[str, Any]]:
    totals: dict[str, dict[str, Any]] = {}
    count_headers = [
        header for header in headers
        if header not in {"屆數", "隊伍", "姓名", "組別", "球員ID"} and header not in derived_headers
    ]
    for row in rows:
        player_id = str(row.get("球員ID") or "")
        if not player_id:
            continue
        if player_id not in totals:
            player = players[player_id]
            totals[player_id] = {
                "屆數": "生涯累計",
                "隊伍": player["currentTeam"] or row.get("隊伍", ""),
                "姓名": player["name"],
                "組別": player["division"] or "歷屆",
                "球員ID": player_id,
                "姓名別名": "、".join(player["aliases"]),
                "_latestSeason": "",
                "_knownHeaders": set(),
            }
            for header in count_headers:
                totals[player_id][header] = 0
        total = totals[player_id]
        if str(row.get("屆數") or "") > str(total["_latestSeason"]):
            total["_latestSeason"] = row.get("屆數", "")
            if not players[player_id]["currentTeam"]:
                total["隊伍"] = row.get("隊伍", "")
        for header in count_headers:
            if row.get(header) not in (None, ""):
                total["_knownHeaders"].add(header)
                total[header] = number(total.get(header)) + number(row.get(header))

    output: list[dict[str, Any]] = []
    for player_id, total in totals.items():
        total.pop("_latestSeason", None)
        known_headers = total.pop("_knownHeaders", set())
        for header in count_headers:
            total[header] = compact_number(number(total.get(header))) if header in known_headers else ""
        if kind == "hitter":
            ab, hits = number(total.get("打數")), number(total.get("安打"))
            walks, hbp, sf = number(total.get("保送")), number(total.get("觸身球")), number(total.get("犧牲飛球"))
            pa, tb, strikeouts = number(total.get("打席")), number(total.get("壘打數")), number(total.get("三振"))
            total["打擊率"] = rate(hits, ab, 3)
            total["上壘率"] = rate(hits + walks + hbp, ab + walks + hbp + sf, 3)
            total["長打率"] = rate(tb, ab, 3)
            total["被三振率"] = rate(strikeouts, pa, 3)
            total["OPS"] = round(number(total["上壘率"]) + number(total["長打率"]), 3)
        else:
            outs = number(total.get("總出局數"))
            earned, runs = number(total.get("自責分")), number(total.get("總失分"))
            strikeouts, walks, hits = number(total.get("三振")), number(total.get("保送")), number(total.get("被安打"))
            total["投球局數"] = innings_from_outs(outs)
            total["非自責分"] = compact_number(max(0, runs - earned))
            total["防禦率"] = round(earned * 27 / outs, 2) if outs else 0.0
            total["K9"] = round(strikeouts * 27 / outs, 2) if outs else 0.0
            total["WHIP"] = round((hits + walks) * 3 / outs, 2) if outs else 0.0
            total["BB9"] = round(walks * 27 / outs, 2) if outs else 0.0
            total["H9"] = round(hits * 27 / outs, 2) if outs else 0.0
        output.append(total)
    return sorted(output, key=lambda item: (item["姓名"], item["球員ID"]))


def main() -> None:
    args = parse_args()
    if not args.source.exists():
        raise SystemExit(f"Source workbook not found: {args.source}")

    workbook = load_workbook(args.source, read_only=True, data_only=True)
    master_rows = table_rows(workbook["永久球員主檔"])
    players: dict[str, dict[str, Any]] = {}
    for row in master_rows:
        player_id = str(row.get("球員ID") or "").strip()
        if not player_id or row.get("資料檢核") != "OK" or row.get("狀態") != "有效":
            continue
        aliases = split_aliases(row.get("姓名別名"))
        old_ids = split_aliases(row.get("舊球員ID"))
        current_team = canonical_team_name(
            row.get("第三屆隊伍") or row.get("第二屆隊伍") or row.get("第一屆隊伍") or ""
        )
        players[player_id] = {
            "id": player_id,
            "name": str(row.get("標準姓名") or "").strip(),
            "aliases": aliases,
            "oldIds": old_ids,
            "teams": {
                "第一屆": canonical_team_name(row.get("第一屆隊伍")),
                "第二屆": canonical_team_name(row.get("第二屆隊伍")),
                "第三屆": canonical_team_name(row.get("第三屆隊伍")),
            },
            "currentTeam": current_team,
            "division": row.get("第三屆組別") or "",
        }

    if len(players) != len({player["id"] for player in players.values()}):
        raise SystemExit("Duplicate permanent player IDs found")
    unnamed = [player_id for player_id, player in players.items() if not player["name"]]
    if unnamed:
        raise SystemExit(f"Players missing names: {', '.join(unnamed[:10])}")

    hitter_rows: list[dict[str, Any]] = []
    for row in table_rows(workbook["歷屆打者成績"]):
        player_id = str(row.get("球員ID") or "").strip()
        if approved_historical_row(row) and player_id in players:
            hitter_rows.append(public_row(row, HITTER_HEADERS, players[player_id], HITTER_ALIASES))

    pitcher_rows: list[dict[str, Any]] = []
    for row in table_rows(workbook["歷屆投手成績"]):
        player_id = str(row.get("球員ID") or "").strip()
        if approved_historical_row(row) and player_id in players:
            pitcher_rows.append(public_row(row, PITCHER_HEADERS, players[player_id]))

    add_third_season_rows(args.hitter_workbook, HITTER_HEADERS, players, hitter_rows, HITTER_ALIASES)
    add_third_season_rows(args.pitcher_workbook, PITCHER_HEADERS, players, pitcher_rows)

    hitter_career = career_rows(hitter_rows, HITTER_HEADERS, players, HITTER_DERIVED, "hitter")
    pitcher_career = career_rows(pitcher_rows, PITCHER_HEADERS, players, PITCHER_DERIVED, "pitcher")
    generated_at = datetime.now(timezone.utc).astimezone().replace(microsecond=0).isoformat()
    source_files = [args.source]
    for optional_source in [args.hitter_workbook, args.pitcher_workbook]:
        if optional_source.exists():
            source_files.append(optional_source)

    payload = {
        "meta": {
            "schemaVersion": 1,
            "generatedAt": generated_at,
            "sourceWorkbook": args.source.name,
            "sourceSha256": sha256(args.source),
            "sourceFiles": [path.name for path in source_files],
            "masterVersion": max((str(row.get("主檔版本") or "") for row in master_rows), default=""),
            "counts": {
                "players": len(players),
                "hitterSeasonRows": len(hitter_rows),
                "pitcherSeasonRows": len(pitcher_rows),
                "hitterCareerRows": len(hitter_career),
                "pitcherCareerRows": len(pitcher_career),
            },
        },
        "players": sorted(players.values(), key=lambda item: (item["name"], item["id"])),
        "datasets": {
            "hitter": {
                "headers": HITTER_HEADERS,
                "rows": hitter_rows,
                "careerRows": hitter_career,
            },
            "pitcher": {
                "headers": PITCHER_HEADERS,
                "rows": pitcher_rows,
                "careerRows": pitcher_career,
            },
        },
    }

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(json.dumps(payload["meta"], ensure_ascii=False, indent=2))
    print(f"Output: {args.output}")


if __name__ == "__main__":
    main()
