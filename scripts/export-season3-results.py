#!/usr/bin/env python3
"""Export the official Season 3 results workbook to website data."""

from __future__ import annotations

import argparse
import json
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


GROUP_IDS = {"黑洞組": "black-hole", "恆星組": "stellar", "跨組交流": "crossover"}
TEAM_IDS = {
    "安穆科技": "anmu-tech",
    "TCW": "tcw",
    "好CHILL": "good-chill",
    "捷創科技": "jtron",
    "聚福運動": "jf-sport",
    "地獄天使": "hell-angels",
    "樂天": "rakuten",
    "台中社會人": "taichung-social",
    "安穆壞壞熊": "bad-bears",
    "創邑": "chungyi",
    "NGU": "ngu",
    "遊牧者": "nomads",
}
FINAL_STATUSES = {"完賽", "沒收"}
VALID_STATUSES = {"未賽", "完賽", "延期", "沒收"}


def iso_date(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    raise ValueError(f"日期欄位格式錯誤：{value!r}")


def read_schedule(path: Path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        teams = {}
        for name, group in workbook["名單與設定"].iter_rows(
            min_row=9, max_row=28, min_col=1, max_col=2, values_only=True
        ):
            if name:
                if name not in TEAM_IDS or group not in GROUP_IDS:
                    raise ValueError(f"官網尚未設定球隊或組別：{name}／{group}")
                teams[name] = group

        if len(teams) != 12:
            raise ValueError(f"正式名單應為 12 隊，目前讀到 {len(teams)} 隊。")

        games = {}
        rows = workbook["賽程排班"].iter_rows(
            min_row=9, max_row=168, min_col=1, max_col=16, values_only=True
        )
        for row in rows:
            game_id, _, day, game_time, slot, group, away, home, *_rest, schedule_status = row
            if not day:
                continue
            if not game_id or game_id in games:
                raise ValueError("賽程場次 ID 空白或重複。")
            games[game_id] = {
                "id": game_id,
                "day": day,
                "time": game_time,
                "slot": slot,
                "group": group,
                "away": away,
                "home": home,
                "venue": row[11],
                "schedule_status": schedule_status,
            }
        if len(games) != 36:
            raise ValueError(f"上半季賽程應為 36 場，目前讀到 {len(games)} 場。")
        return games, teams
    finally:
        workbook.close()


def read_results(path: Path):
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = workbook["賽果登錄"]
        expected = [
            "場次 ID", "日期", "時間", "場次", "組別", "客隊／先攻", "主隊／後攻",
            "賽果狀態", "客隊得分", "主隊得分", "沒收方", "確認日期", "備註", "資料核對",
        ]
        header = list(next(sheet.iter_rows(min_row=8, max_row=8, min_col=1, max_col=14, values_only=True)))
        if header != expected:
            raise ValueError("賽果登錄欄位已變更，停止匯入。")

        results = {}
        for row in sheet.iter_rows(min_row=9, max_row=168, min_col=1, max_col=13, values_only=True):
            if not row[0]:
                continue
            if row[0] in results:
                raise ValueError(f"賽果場次 ID 重複：{row[0]}")
            results[row[0]] = {
                "id": row[0], "day": row[1], "time": row[2], "slot": row[3], "group": row[4],
                "away": row[5], "home": row[6], "status": row[7] or "未賽",
                "away_score": row[8], "home_score": row[9], "forfeiter": row[10],
                "confirmed_at": row[11], "note": row[12],
            }
        return results
    finally:
        workbook.close()


def validate(games, teams, results):
    accepted = []
    for game_id, game in games.items():
        result = results.get(game_id)
        if result is None:
            continue
        status = result["status"]
        if status not in VALID_STATUSES:
            raise ValueError(f"{game_id} 賽果狀態無效：{status}")
        if status not in FINAL_STATUSES:
            if any(value not in (None, "") for value in (
                result["away_score"], result["home_score"], result["forfeiter"], result["confirmed_at"]
            )):
                raise ValueError(f"{game_id} 尚未完賽卻已有比分或確認日期。")
            continue

        current = (game["day"], game["time"], game["slot"], game["group"], game["away"], game["home"])
        recorded = (result["day"], result["time"], result["slot"], result["group"], result["away"], result["home"])
        if current != recorded:
            raise ValueError(f"{game_id} 賽程與賽果的日期、場次、組別或球隊不同。")
        if not all(isinstance(score, int) and not isinstance(score, bool) and 0 <= score <= 99
                   for score in (result["away_score"], result["home_score"])):
            raise ValueError(f"{game_id} 比分必須是 0 至 99 的整數。")
        if not isinstance(result["confirmed_at"], datetime) or result["confirmed_at"].date() < game["day"].date():
            raise ValueError(f"{game_id} 確認日期不得早於比賽日。")
        if status == "沒收":
            expected = (0, 9) if result["forfeiter"] == "客隊" else (9, 0) if result["forfeiter"] == "主隊" else None
            if expected is None or (result["away_score"], result["home_score"]) != expected:
                raise ValueError(f"{game_id} 沒收方或沒收比分不正確。")
        elif result["forfeiter"]:
            raise ValueError(f"{game_id} 一般完賽不可設定沒收方。")
        if game["away"] not in teams or game["home"] not in teams:
            raise ValueError(f"{game_id} 球隊不在第三屆正式名單。")
        accepted.append((game, result))

    extra = [game_id for game_id, result in results.items()
             if game_id not in games and result["status"] in FINAL_STATUSES]
    if extra:
        raise ValueError(f"賽果包含已移除的正式場次：{', '.join(extra)}")
    return accepted


def calculate(teams, accepted):
    stats = {
        team: {
            "team": team, "group": group, "played": 0, "wins": 0, "draws": 0, "losses": 0,
            "points": 0, "scored": 0, "allowed": 0, "difference": 0,
            "winningPercentage": None, "ratio": None, "recent": [], "rank": None,
        }
        for team, group in teams.items()
    }
    history = defaultdict(list)
    for game, result in accepted:
        away_score, home_score = result["away_score"], result["home_score"]
        pairings = (
            (game["away"], away_score, home_score, game["home"]),
            (game["home"], home_score, away_score, game["away"]),
        )
        for team, scored, allowed, opponent in pairings:
            row = stats[team]
            row["played"] += 1
            row["scored"] += scored
            row["allowed"] += allowed
            outcome = "勝" if scored > allowed else "敗" if scored < allowed else "和"
            row["wins"] += outcome == "勝"
            row["draws"] += outcome == "和"
            row["losses"] += outcome == "敗"
            row["points"] += 2 if outcome == "勝" else 1 if outcome == "和" else 0
            history[team].append((game["day"], game["time"], game["slot"], game["id"], outcome, opponent))

    for team, row in stats.items():
        row["difference"] = row["scored"] - row["allowed"]
        if row["played"]:
            row["winningPercentage"] = row["wins"] / row["played"]
            row["recent"] = [item[4] for item in sorted(history[team], reverse=True)[:5]]
        if row["allowed"]:
            row["ratio"] = (row["scored"] + row["allowed"]) / row["allowed"]

    def head_to_head(team, peers):
        return sum(2 if outcome == "勝" else 1 if outcome == "和" else 0
                   for *_, outcome, opponent in history[team] if opponent in peers)

    standings = {}
    for group in GROUP_IDS:
        active = [row for row in stats.values() if row["group"] == group and row["played"]]
        inactive = [row for row in stats.values() if row["group"] == group and not row["played"]]
        points_groups = defaultdict(list)
        for row in active:
            points_groups[row["points"]].append(row)
        ordered = []
        for points in sorted(points_groups, reverse=True):
            peers = points_groups[points]
            peer_names = {row["team"] for row in peers}
            for row in peers:
                row["h2h"] = head_to_head(row["team"], peer_names)
            peers.sort(key=lambda row: (
                -row["h2h"], -(row["ratio"] if row["ratio"] is not None else -1),
                row["allowed"], -row["scored"], row["team"],
            ))
            if len(peers) > 1 and any(row["ratio"] is None for row in peers):
                tied_by_head_to_head = defaultdict(list)
                for row in peers:
                    tied_by_head_to_head[row["h2h"]].append(row)
                for subset in tied_by_head_to_head.values():
                    if len(subset) > 1 and any(row["ratio"] is None for row in subset):
                        for row in subset:
                            row["unresolved"] = True
            for index, row in enumerate(peers):
                comparison = (row["h2h"], row["ratio"], row["allowed"], row["scored"])
                previous = peers[index - 1] if index else None
                previous_comparison = (
                    previous["h2h"], previous["ratio"], previous["allowed"], previous["scored"]
                ) if previous else None
                if row.get("unresolved"):
                    row["rank"] = "待裁定"
                elif previous and comparison == previous_comparison:
                    row["rank"] = previous["rank"]
                else:
                    row["rank"] = len(ordered) + index + 1
            ordered.extend(peers)
        ordered.extend(sorted(inactive, key=lambda row: TEAM_IDS[row["team"]]))
        standings[GROUP_IDS[group]] = ordered
    return standings


def build_payload(schedule_path: Path, results_path: Path):
    games, teams = read_schedule(schedule_path)
    results = read_results(results_path)
    accepted = validate(games, teams, results)
    standings = calculate(teams, accepted)
    confirmed_dates = [result["confirmed_at"] for _, result in accepted]
    game_results = {}
    for game_id, result in results.items():
        if result["status"] == "未賽":
            continue
        entry = {"status": result["status"]}
        if result["status"] in FINAL_STATUSES:
            entry.update({
                "awayScore": result["away_score"],
                "homeScore": result["home_score"],
                "confirmedAt": iso_date(result["confirmed_at"]),
            })
        game_results[game_id] = entry

    clean_standings = {}
    for group_id, rows in standings.items():
        clean_standings[group_id] = [{
            "teamId": TEAM_IDS[row["team"]],
            "rank": row["rank"],
            "played": row["played"],
            "wins": row["wins"],
            "draws": row["draws"],
            "losses": row["losses"],
            "points": row["points"],
            "scored": row["scored"],
            "allowed": row["allowed"],
            "difference": row["difference"],
            "winningPercentage": row["winningPercentage"],
            "recent": row["recent"],
        } for row in rows]

    source_date = datetime.fromtimestamp(max(schedule_path.stat().st_mtime, results_path.stat().st_mtime)).date().isoformat()
    return {
        "season": 3,
        "phase": "first-half",
        "state": "complete" if len(accepted) == len(games) else "active" if accepted else "preseason",
        "sourceUpdatedAt": source_date,
        "standingsUpdatedAt": max(confirmed_dates).date().isoformat() if confirmed_dates else None,
        "confirmedGames": len(accepted),
        "gameResults": game_results,
        "standings": clean_standings,
        "rules": {
            "winPoints": 2,
            "drawPoints": 1,
            "lossPoints": 0,
            "crossoverCounts": True,
            "note": "同分時依聯盟規章比較；待裁定名次以聯盟正式公告為準。",
        },
    }


def build_schedule_payload(schedule_path: Path):
    games, _teams = read_schedule(schedule_path)
    crossover_rank = 0
    clean_games = []
    for game in sorted(games.values(), key=lambda item: (item["day"], item["time"])):
        if game["group"] not in GROUP_IDS:
            raise ValueError(f"賽程組別無法轉換：{game['group']}")
        division = GROUP_IDS[game["group"]]
        entry = {
            "id": game["id"],
            "date": game["day"].date().isoformat(),
            "time": game["time"].strftime("%H:%M"),
            "slot": int(game["slot"]),
            "division": division,
            "venue": game["venue"] or "大肚溪A球場",
        }
        if division == "crossover":
            crossover_rank += 1
            entry["leftSeed"] = f"黑洞組第 {crossover_rank} 名"
            entry["rightSeed"] = f"恆星組第 {crossover_rank} 名"
        else:
            if game["away"] not in TEAM_IDS or game["home"] not in TEAM_IDS:
                raise ValueError(f"{game['id']} 球隊無法轉換為官網 ID。")
            entry["away"] = TEAM_IDS[game["away"]]
            entry["home"] = TEAM_IDS[game["home"]]
        clean_games.append(entry)

    return {
        "title": "第三屆上半季賽程",
        "venue": clean_games[0]["venue"] if clean_games else "大肚溪A球場",
        "updatedAt": datetime.fromtimestamp(schedule_path.stat().st_mtime).date().isoformat(),
        "notice": "如有賽程異動，以聯盟最新公告為準。",
        "games": clean_games,
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--schedule", type=Path, required=True)
    parser.add_argument("--results", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--schedule-output", type=Path)
    args = parser.parse_args()
    payload = build_payload(args.schedule, args.results)
    output = "(function () {\n  \"use strict\";\n\n  window.CBL_SEASON3_RESULTS = "
    output += json.dumps(payload, ensure_ascii=False, indent=2)
    output += ";\n})();\n"
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(output, encoding="utf-8")
    report = {"confirmedGames": payload["confirmedGames"], "state": payload["state"], "output": str(args.output)}
    if args.schedule_output:
        schedule_payload = build_schedule_payload(args.schedule)
        schedule_output = "(function () {\n  \"use strict\";\n\n  window.CBL_SEASON3_SCHEDULE = "
        schedule_output += json.dumps(schedule_payload, ensure_ascii=False, indent=2)
        schedule_output += ";\n})();\n"
        args.schedule_output.parent.mkdir(parents=True, exist_ok=True)
        args.schedule_output.write_text(schedule_output, encoding="utf-8")
        report["scheduleOutput"] = str(args.schedule_output)
        report["scheduleGames"] = len(schedule_payload["games"])
    print(json.dumps(report, ensure_ascii=False))


if __name__ == "__main__":
    main()
